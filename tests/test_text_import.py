from io import BytesIO

from docx import Document

from metrotrance.services.text_import import import_text


def test_import_utf8_text():
    text = import_text("practice.txt", "Первая строка.\n\nВторая строка.".encode("utf-8"))
    assert "Первая" in text
    assert "Вторая" in text


def test_import_docx_preserves_paragraph_boundaries():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Первый абзац спокойной практики.")
    document.add_paragraph("Второй абзац с новой смысловой сценой.")
    document.save(buffer)
    text = import_text("practice.docx", buffer.getvalue())
    assert "\n\n" in text
    assert "Второй абзац" in text
